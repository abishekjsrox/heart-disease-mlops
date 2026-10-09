"""Builds the final report (reports/MLOps_Assignment1_Report.docx).

Numbers, tables and figures are read from the project outputs, and screenshots are
picked up from screenshots/. Any screenshot that is still missing shows up as a red
"ADD SCREENSHOT" note, so re-run this script after capturing them.

Usage:  python scripts/build_report.py
"""
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt, RGBColor

from heartrisk import config

ROOT = config.ROOT_DIR
FIG = config.FIGURES_DIR
SHOTS = ROOT / "screenshots"
CFG = json.loads((ROOT / "reports" / "report_config.json").read_text())
META = json.loads(config.MODEL_META_FILE.read_text())
EDA = json.loads((config.REPORTS_DIR / "eda_summary.json").read_text())
CMP = pd.read_csv(config.REPORTS_DIR / "model_comparison.csv", index_col=0)
OUT = config.REPORTS_DIR / "MLOps_Assignment1_Report.docx"
missing = []

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Cm(2.0))
sec.top_margin = sec.bottom_margin = Cm(1.8)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(4)
for lvl, size in ((1, 14), (2, 11.5)):
    st = doc.styles[f"Heading {lvl}"]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    st.paragraph_format.space_before = Pt(10 if lvl == 1 else 6)
    st.paragraph_format.space_after = Pt(3)


def h(text, lvl=1):
    doc.add_heading(text, level=lvl)


def p(text, bold_prefix=None):
    para = doc.add_paragraph()
    if bold_prefix:
        para.add_run(bold_prefix).bold = True
    para.add_run(text)
    return para


def bullets(items):
    for it in items:
        para = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            para.add_run(it[0]).bold = True
            para.add_run(it[1])
        else:
            para.add_run(it)


def code(text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Cm(0.4)
    run = para.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    return para


def table(rows, header, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, col in enumerate(header):
        t.rows[0].cells[i].text = str(col)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    for row in t.rows:
        for i, c in enumerate(row.cells):
            for para in c.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9)
            if widths:
                c.width = Cm(widths[i])
    doc.add_paragraph()


def image(path: Path, caption: str, width_cm=16.0, hint: str = ""):
    if path.exists():
        doc.add_picture(str(path), width=Cm(width_cm))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        missing.append(path.name)
        para = doc.add_paragraph()
        run = para.add_run(f"[ADD SCREENSHOT: screenshots/{path.name}]  {hint}")
        run.bold = True
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


m = META["metrics"]
champ = META["model_name"].replace("_", " ")

# ------------------------------------------------------------------ title ---
for _ in range(5):
    doc.add_paragraph()
for text, size, bold in (("Machine Learning Operations (AIMLCZG523)", 16, True),
                         ("Assignment 1 - End-to-End MLOps Pipeline", 14, True),
                         ("Heart Disease Risk Prediction API", 13, False)):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = para.add_run(text)
    r.font.size, r.bold = Pt(size), bold
doc.add_paragraph()
table([["Name", CFG["student_name"]], ["BITS ID", CFG["student_id"]],
       ["Programme", "M.Tech AI/ML (WILP), Semester 3"],
       ["Code repository", CFG["repo_url"]],
       ["Pipeline video", CFG.get("video_url", "")],
       ["Submission date", CFG.get("submission_date", "")]],
      ["Item", "Details"], widths=[4.5, 11.5])
page_break()

# ------------------------------------------------------------ 1 overview ---
h("1. Problem and approach")
p("The goal is to predict whether a patient has heart disease from 13 routine clinical "
  "attributes (UCI Heart Disease, Cleveland subset) and to run the model as a monitored "
  "API, built the way it would be in production: reproducible data steps, tracked "
  "experiments, automated tests and builds, a container, a Kubernetes deployment and "
  "basic monitoring.")
p("I treated the model as a screening aid. In that setting, a sick patient sent home as "
  "healthy (false negative) costs much more than an extra test for a healthy patient, so "
  "recall on the disease class mattered more to me than raw accuracy. This choice shows "
  "up in model selection, the decision threshold and the monitoring dashboard.")
image(FIG / "architecture.png", "Figure 1: End-to-end architecture", 15)

h("1.1 Repository layout", 2)
code("""heart-disease-mlops/
  src/heartrisk/        config, data, features, train, predict, api (FastAPI), eda
  scripts/              download_data.py, run_eda.py, load_test.py, build_report.py
  notebooks/            01_eda, 02_training, 03_inference (executed)
  data/raw|processed/   raw UCI file and cleaned CSV
  models/               model.joblib + model_meta.json (champion)
  tests/                pytest suite (data, features, model, api, eda)
  .github/workflows/    ci.yml (lint -> test -> train -> docker smoke test)
  Dockerfile, docker-compose.yml, k8s/, monitoring/ (Prometheus + Grafana)
  screenshots/, reports/""")

h("1.2 Setup and run instructions", 2)
code("""git clone <repo> && cd heart-disease-mlops
python3.11 -m venv .venv && source .venv/bin/activate
pip install -r requirements.txt          # also installs the package (-e .)
python scripts/download_data.py          # UCI download -> data/raw, data/processed
python scripts/run_eda.py                # figures -> reports/figures
python -m heartrisk.train                # tuning + MLflow + models/model.joblib
pytest -v                                # unit tests
mlflow ui --backend-store-uri sqlite:///mlflow.db --port 5000
docker build -t heart-risk-api:1.0.0 . && docker run -p 8000:8000 heart-risk-api:1.0.0
kubectl apply -f k8s/                     # Docker Desktop Kubernetes / minikube
docker compose up -d                     # API + Prometheus :9090 + Grafana :3000""")

# --------------------------------------------------------------- 2 data ---
h("2. Data acquisition and EDA")
p("scripts/download_data.py pulls processed.cleveland.data from the UCI repository. If UCI "
  "is unreachable it falls back to a public mirror of the same 303 records and converts it "
  "to the UCI coding, so the rest of the pipeline never changes. The raw file is saved "
  "untouched in data/raw; cleaning writes data/processed/heart_clean.csv.")
bullets([
    ("Missing values: ", f"{EDA['missing_values']} - the raw file uses '?'. They are kept as "
     "NaN in the clean file and imputed inside the model pipeline, so the imputer only sees "
     "training folds (no leakage) and the API applies the identical step."),
    ("Validation: ", "codes outside the documented sets (e.g. thal not in 3/6/7) and values "
     "outside plausible ranges are turned into missing values rather than silently kept."),
    ("Target: ", "num (0-4) is converted to binary - 0 = no disease, 1-4 = disease."),
    ("Class balance: ", f"{EDA['class_counts']} - mildly imbalanced, so accuracy is still "
     "meaningful but recall and ROC-AUC are reported alongside it."),
])
image(FIG / "02_histograms.png", "Figure 2: Continuous features by outcome", 14)
image(FIG / "03_correlation_heatmap.png", "Figure 3: Correlation heatmap", 9.5)
p("Strongest associations with the target: " + ", ".join(
    f"{k} ({v})" for k, v in EDA["top_correlations_with_target"].items()) + ". Patients with "
  "disease reach a lower peak heart rate and show more ST depression. Asymptomatic chest pain "
  "(cp = 4) has the highest disease rate (about 73%), which looks odd but is well known for "
  "this dataset. restecg = 1 has only 4 patients, which is why the encoder ignores unseen "
  "categories instead of failing.")
image(FIG / "04_categorical_disease_rate.png", "Figure 4: Disease rate per category", 13.5)

# ----------------------------------------------------------- 3 modelling ---
h("3. Feature engineering and model development")
p("All preprocessing lives in one sklearn Pipeline that is saved together with the model:")
bullets([
    ("Engineered feature: ", "hr_reserve_pct = thalach / (220 - age), the share of the "
     "age-predicted maximum heart rate reached in the stress test. It combines two raw "
     "columns into one clinically meaningful signal."),
    ("Numeric (age, trestbps, chol, thalach, oldpeak, ca, hr_reserve_pct): ",
     "median imputation + standard scaling."),
    ("Binary (sex, fbs, exang): ", "most-frequent imputation."),
    ("Categorical (cp, restecg, slope, thal): ",
     "most-frequent imputation + one-hot encoding (handle_unknown = ignore)."),
])
p("I compared three model families. Data was split 80/20 (stratified, seed 42). Each "
  "model was tuned with GridSearchCV on the 80% part using 5-fold stratified CV and ROC-AUC "
  "as the score. The champion is the best mean CV ROC-AUC, so the hold-out set is never "
  "used for any choice and only gives the final, unbiased numbers.")
table([["Logistic Regression", "C, L1/L2 penalty, class_weight", "20"],
       ["Random Forest", "n_estimators, max_depth, min_samples_leaf, max_features", "36"],
       ["Gradient Boosting", "n_estimators, learning_rate, max_depth, subsample", "16"]],
      ["Model", "Tuned hyper-parameters", "Combinations"], widths=[4, 9.5, 2.5])
rows = [[k.replace("_", " "), f"{r.cv_roc_auc:.3f} +/- {r.cv_roc_auc_std:.3f}",
         f"{r.cv_accuracy:.3f}", f"{r.cv_precision:.3f}", f"{r.cv_recall:.3f}",
         f"{r.test_roc_auc:.3f}", f"{r.test_recall:.3f}", f"{r.threshold:.2f}"]
        for k, r in CMP.iterrows()]
table(rows, ["Model", "CV ROC-AUC", "CV acc.", "CV prec.", "CV recall",
             "Test AUC", "Test recall*", "Threshold"])
p(f"The champion is {champ} (best parameters: {META['best_params']}). With only 242 "
  "training rows, the simpler linear model generalised as well as or better than the tree "
  "ensembles, and it is also the easiest to explain to a clinician. CV metrics use the "
  "default 0.5 cut-off. *Test metrics use the tuned threshold.")
h("3.1 Decision threshold", 2)
p(f"At the default 0.5 cut-off the champion's cross-validated recall was only "
  f"{m['cv_recall']:.2f}. I therefore picked the threshold from out-of-fold training "
  f"predictions as the highest value that still reaches {META['target_recall']:.0%} recall: "
  f"{META['decision_threshold']}. On the untouched hold-out set this gives recall "
  f"{m['test_recall']:.3f}, precision {m['test_precision']:.3f} and accuracy "
  f"{m['test_accuracy']:.3f} (ROC-AUC {m['test_roc_auc']:.3f}). The trade-off is more false "
  "alarms, which is acceptable for screening. The threshold is stored in model_meta.json "
  "and applied by the API, so it is versioned together with the model.")
image(FIG / f"confusion_matrix_{META['model_name']}.png",
      "Figure 5: Hold-out confusion matrix at the tuned threshold", 7.5)
image(FIG / f"feature_importance_{META['model_name']}.png",
      "Figure 6: Largest model coefficients / importances", 10)

# ------------------------------------------------------------ 4 tracking ---
h("4. Experiment tracking (MLflow)")
p("Each model family is one MLflow run in the experiment heart-disease-risk, backed by a "
  "local SQLite store (mlflow.db) with artifacts under mlruns/. Per run I log:")
bullets([
    "parameters: best grid values, CV folds, split sizes, grid size, decision threshold",
    "metrics: CV accuracy / precision / recall / F1 / ROC-AUC (+ std), hold-out metrics at "
    "both the tuned threshold and 0.5",
    "artifacts: confusion matrix, ROC curve, feature importance plot, full grid-search "
    "results CSV, and the model itself with input signature and example",
    "tags: model family, git commit SHA, dataset id, champion flag",
])
p("The champion is registered as heart-disease-classifier with the alias 'champion', so a "
  "deployment can always refer to the approved version instead of a run id.")
image(SHOTS / "mlflow_03_compare_metrics.png", "Figure 7: MLflow - comparing the three runs",
      12.5)
image(SHOTS / "mlflow_04_champion_run.png", "Figure 8: Champion run - metrics and tags", 12.5)
image(SHOTS / "mlflow_05_artifacts.png", "Figure 9: Logged artifacts (ROC curve)", 12.5)

# ------------------------------------------------------- 5 packaging -------
h("5. Model packaging and reproducibility")
bullets([
    ("Single artifact: ", "models/model.joblib holds preprocessing + model, so training and "
     "serving cannot drift apart. The same model is also saved in MLflow format."),
    ("Metadata: ", "models/model_meta.json records model version (package version + git "
     "SHA), MLflow run id, registry version, threshold, metrics, sklearn and Python versions."),
    ("Pinned environment: ", "requirements.txt pins every library (Python 3.11). The serving "
     "image uses requirements-serve.txt, a subset with the same versions as training."),
    ("Determinism: ", "fixed seed for the split, CV folds and models; data download is "
     "scripted and the cleaned CSV is committed."),
])
code(json.dumps({k: META[k] for k in ("model_name", "model_version", "decision_threshold",
                                      "registry_version", "sklearn_version")}, indent=1))

# ------------------------------------------------------------- 6 CI/CD ---
h("6. CI/CD pipeline and automated testing")
p(f"The test suite has {CFG.get('test_count', 'all')} pytest tests across five files: "
  "parsing of both data sources, cleaning rules (invalid codes, ranges, duplicates, target), "
  "the engineered feature, imputation and scaling, unseen categories, threshold selection, a "
  "minimum quality bar for the saved model, and API behaviour (valid prediction, 422 on "
  "bad input, batch endpoint, metrics endpoint).")
p("GitHub Actions (.github/workflows/ci.yml) runs on every push and pull request to main. "
  "Each job needs the previous one, so any lint error, failed test or broken build stops "
  "the pipeline with the log of the failing step:")
table([["1. Lint", "flake8 on src, tests, scripts", "-"],
       ["2. Unit tests", "download data, quick training, pytest with coverage",
        "junit.xml, coverage.xml"],
       ["3. Train", "full training, EDA, quality gate (tests/test_model.py)",
        "model, mlruns, mlflow.db, reports, training.log"],
       ["4. Docker", "build image, run it, wait for /ready, check /predict, 422 and /metrics",
        "container.log"]],
      ["Job", "What it does", "Artifacts"], widths=[2.8, 8.7, 4.5])
image(SHOTS / "ci_01_workflow_run.png", "Figure 10: Successful GitHub Actions run", 12.5,
      "GitHub repo -> Actions -> latest run (shows the 4 green jobs)")
image(SHOTS / "ci_02_artifacts.png", "Figure 11: Workflow artifacts and job summary", 12.5,
      "Bottom of the same run page: artifacts list + model comparison summary")

# ---------------------------------------------------------- 7 container ---
h("7. Containerisation")
p("The API is a FastAPI app (src/heartrisk/api.py). POST /predict takes one patient as JSON "
  "and returns the prediction, a readable label, the disease probability, the confidence "
  "in the returned class, a risk band, the threshold and the model version. Input is "
  "validated with Pydantic (allowed codes and clinical ranges), so bad requests get HTTP "
  "422 instead of a silent wrong answer. ca and thal may be null and are imputed. Other "
  "endpoints: /predict/batch, /health (liveness), /ready (model loaded), /model-info, "
  "/metrics and /docs.")
p("The Dockerfile uses python:3.11-slim, installs only the serving requirements in a cached "
  "layer, runs as a non-root user and has a HEALTHCHECK. One uvicorn worker runs per "
  "container; scaling is done with Kubernetes replicas, which also keeps the Prometheus "
  "counters per process simple.")
code("""$ curl -X POST localhost:8000/predict -H 'Content-Type: application/json' \\
       -d @samples/high_risk.json
{"prediction":1,"label":"heart disease risk","probability_disease":0.9932,
 "confidence":0.9932,"risk_band":"high","threshold":0.3,"model_version":"1.0.0+..."}""")
image(SHOTS / "docker_01_build_run.png", "Figure 12: Image built and container running "
      "locally", 12.5, "Terminal: docker build ... then docker ps")
image(SHOTS / "docker_02_predict.png", "Figure 13: /predict on the local container", 12.5,
      "Terminal: curl /predict with samples/high_risk.json and low_risk.json")
image(SHOTS / "api_01_swagger.png", "Figure 14: Interactive API docs (/docs)", 12.5)

# --------------------------------------------------------------- 8 k8s ---
h("8. Production deployment (Kubernetes)")
p(f"I deployed to {CFG.get('k8s_platform', 'Docker Desktop Kubernetes')} using the manifests "
  "in k8s/: a namespace, a Deployment with 2 replicas, and a LoadBalancer Service on port "
  "80. The Deployment uses a readiness probe on /ready (no traffic until the model is "
  "loaded), a liveness probe on /health, CPU and memory requests and limits, a non-root "
  "security context and a rolling update with maxUnavailable 0, so a new model version "
  "rolls out without downtime. An optional nginx Ingress (heart-api.local) is included.")
code("""docker build -t heart-risk-api:1.0.0 .
kubectl apply -f k8s/namespace.yaml -f k8s/deployment.yaml -f k8s/service.yaml
kubectl -n heart-risk get pods,svc
curl -X POST http://localhost/predict -H 'Content-Type: application/json' \\
     -d @samples/high_risk.json""")
image(SHOTS / "k8s_01_pods_svc.png", "Figure 15: Pods running and LoadBalancer service",
      12.5, "Terminal: kubectl -n heart-risk get pods,svc -o wide")
image(SHOTS / "k8s_02_predict_lb.png", "Figure 16: Prediction through the LoadBalancer",
      12.5, "Terminal: curl http://localhost/predict ... and curl http://localhost/ready")

# ------------------------------------------------------- 9 monitoring ---
h("9. Monitoring and logging")
p("Every request is logged as one JSON line (request id, method, path, status, latency, "
  "client) and every prediction logs the outcome and probability. Logs go to stdout, so "
  "kubectl logs and any log collector can read them. The API also exposes Prometheus "
  "metrics:")
bullets([
    "api_requests_total (by endpoint and status) and api_request_latency_seconds histogram",
    "model_predictions_total by outcome and model_prediction_probability histogram - a "
    "shift in these is an early, label-free sign of data drift",
    "api_invalid_requests_total (validation failures) and model_info (version in service)",
])
p("docker compose starts the API with Prometheus (scraping every 5 s, both the compose "
  "container and the Kubernetes LoadBalancer) and Grafana with a provisioned dashboard: "
  "request rate, p50/p95 latency, error rate, prediction mix, share of disease predictions "
  "and mean predicted probability. scripts/load_test.py sends realistic traffic, including "
  "about 5% invalid requests, to populate it.")
image(SHOTS / "grafana_01_dashboard.png", "Figure 17: Grafana dashboard under load", 12.5,
      "http://localhost:3000 -> Dashboards -> Heart Risk -> Heart Risk API (after load test)")
image(SHOTS / "prometheus_01_targets.png", "Figure 18: Prometheus scrape targets", 12.5,
      "http://localhost:9090/targets")
image(SHOTS / "logs_01_json_requests.png", "Figure 19: Structured request logs", 12.5,
      "Terminal: kubectl -n heart-risk logs deploy/heart-api --tail=15")

# --------------------------------------------------------- 10 closing ---
h("10. Conclusion and next steps")
p(f"The final system takes raw UCI data to a monitored API on Kubernetes with every step "
  f"automated and reproducible. The {champ} champion reaches ROC-AUC {m['test_roc_auc']:.3f} "
  f"and recall {m['test_recall']:.3f} on unseen patients.")
p("Limitations: 303 patients from one hospital is a small and dated sample, so the "
  "hold-out numbers have wide uncertainty (61 test rows) and the model should not be used "
  "clinically without validation on local data. Things I would add next: push the image to "
  "a registry and deploy from CI to a managed cluster (EKS/GKE), Prometheus alert rules on "
  "latency, error rate and prediction-mix drift, a scheduled retraining job, and a proper "
  "drift report comparing live inputs with the training data.")
h("Links", 2)
bullets([("Code repository: ", CFG["repo_url"]),
         ("Pipeline video: ", CFG.get("video_url", "")),
         ("API access: ", CFG.get("api_access",
                                  "local - see README, section 'Run the API'"))])

doc.save(OUT)
print(f"saved {OUT}")
if missing:
    print(f"{len(missing)} screenshot(s) still missing:", ", ".join(missing))
else:
    print("all screenshots present")
